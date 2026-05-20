from __future__ import annotations

from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import Any

from ..config import (
    BASE_DIR,
    MLR_GUIDELINES_DIR,
    MLR_LEGAL_GUIDELINE_PATH,
    MLR_MARKETING_GUIDELINE_PATH,
    MLR_REGULATORY_GUIDELINE_PATH,
)
from ..utils import clean_text


GUIDELINE_FILENAMES = {
    "marketing": "Marketing Guidelines - GENERIC.docx",
    "legal": "Legal Guidelines - GENERIC.docx",
    "regulatory": "Regulatory Guidelines - GENERIC.docx",
}


def _candidate_guideline_dirs() -> list[Path]:
    candidates = []
    if MLR_GUIDELINES_DIR:
        candidates.append(MLR_GUIDELINES_DIR)
    candidates.extend(
        [
            BASE_DIR,
            BASE_DIR.parent,
            Path.cwd(),
            Path.home() / "OneDrive" / "Desktop" / "dataset_ai_mapping",
            Path.home() / "Desktop" / "dataset_ai_mapping",
        ]
    )
    unique = []
    seen = set()
    for candidate in candidates:
        resolved = candidate.expanduser()
        key = str(resolved).lower()
        if key not in seen:
            seen.add(key)
            unique.append(resolved)
    return unique


def _resolve_guideline_paths() -> dict[str, Path]:
    configured = {
        "marketing": MLR_MARKETING_GUIDELINE_PATH,
        "legal": MLR_LEGAL_GUIDELINE_PATH,
        "regulatory": MLR_REGULATORY_GUIDELINE_PATH,
    }
    paths: dict[str, Path] = {}
    missing = []
    for kind, filename in GUIDELINE_FILENAMES.items():
        explicit_path = configured.get(kind)
        if explicit_path and explicit_path.exists():
            paths[kind] = explicit_path
            continue
        for folder in _candidate_guideline_dirs():
            candidate = folder / filename
            if candidate.exists():
                paths[kind] = candidate
                break
        if kind not in paths:
            missing.append(filename)
    if missing:
        raise FileNotFoundError(
            "MLR guideline file(s) not found: "
            + ", ".join(missing)
            + ". Configure MLR_GUIDELINES_DIR or the individual MLR_*_GUIDELINE_PATH environment values."
        )
    return paths


def _read_docx_text(path: Path) -> str:
    try:
        from docx import Document

        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        return clean_text("\n".join(parts))
    except Exception:
        pass

    with zipfile.ZipFile(path) as archive:
        xml_data = archive.read("word/document.xml")
    root = ET.fromstring(xml_data)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = []
    for paragraph in root.findall(".//w:p", namespace):
        texts = [node.text or "" for node in paragraph.findall(".//w:t", namespace)]
        value = clean_text("".join(texts))
        if value:
            paragraphs.append(value)
    return clean_text("\n".join(paragraphs))


def _guideline_items(text: str) -> list[str]:
    items = []
    seen = set()
    for raw_line in re.split(r"\n+|(?<=[.!?])\s+", clean_text(text)):
        line = clean_text(re.sub(r"^[\-\u2022*0-9.)\s]+", "", raw_line)).strip(" :;")
        if len(line) < 18:
            continue
        normalized = re.sub(r"\W+", " ", line.lower()).strip()
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        items.append(line)
        if len(items) >= 40:
            break
    return items


def _terms(text: str) -> set[str]:
    stopwords = {
        "about",
        "after",
        "against",
        "also",
        "and",
        "are",
        "based",
        "been",
        "being",
        "can",
        "claim",
        "claims",
        "content",
        "document",
        "for",
        "from",
        "guideline",
        "guidelines",
        "has",
        "have",
        "into",
        "may",
        "must",
        "not",
        "only",
        "should",
        "that",
        "the",
        "their",
        "this",
        "with",
    }
    return {
        token
        for token in re.findall(r"[a-z0-9]+", text.lower())
        if len(token) > 3 and token not in stopwords
    }


def _review_item(item: str, document_terms: set[str]) -> dict[str, Any]:
    item_terms = _terms(item)
    if not item_terms:
        overlap = 0.0
    else:
        overlap = len(item_terms & document_terms) / len(item_terms)
    satisfied = overlap >= 0.22
    return {
        "guideline": item,
        "status": "Satisfied" if satisfied else "Non Satisfied",
        "confidence": round(overlap, 3),
        "reason": (
            "Relevant content appears in the uploaded material."
            if satisfied
            else "The uploaded material does not show enough matching evidence for this guideline."
        ),
    }


def _document_text_from_payload(payload: dict[str, Any]) -> str:
    result = payload.get("result") if isinstance(payload.get("result"), dict) else payload
    metadata = result.get("metadata") or {}
    source_units = metadata.get("source_units") or []
    parts = [
        result.get("full_transcript"),
        result.get("transcript"),
        result.get("content_summary"),
        result.get("summary"),
        result.get("description_of_image_video"),
        metadata.get("full_extracted_text"),
    ]
    for unit in source_units:
        if isinstance(unit, dict):
            parts.append(unit.get("text"))
            parts.append(unit.get("description"))
    for match in result.get("key_message_matches") or []:
        if isinstance(match, dict):
            parts.extend([match.get("brand_product"), match.get("key_message"), match.get("description_of_image_video")])
    return clean_text("\n".join(str(part) for part in parts if part))


def run_mlr_review(payload: dict[str, Any]) -> dict[str, Any]:
    document_text = _document_text_from_payload(payload)
    if not document_text:
        raise ValueError("Run analysis before MLR review so extracted content is available.")

    guideline_paths = _resolve_guideline_paths()
    document_terms = _terms(document_text)
    sections = []
    all_items = []
    for kind, path in guideline_paths.items():
        guideline_text = _read_docx_text(path)
        items = _guideline_items(guideline_text)
        reviewed_items = [_review_item(item, document_terms) for item in items]
        sections.append(
            {
                "name": kind.title(),
                "file_name": path.name,
                "file_path": str(path),
                "satisfied": [item for item in reviewed_items if item["status"] == "Satisfied"],
                "not_satisfied": [item for item in reviewed_items if item["status"] == "Non Satisfied"],
            }
        )
        all_items.extend(reviewed_items)

    satisfied_count = sum(1 for item in all_items if item["status"] == "Satisfied")
    not_satisfied_count = len(all_items) - satisfied_count
    if not_satisfied_count == 0 and satisfied_count:
        status = "Compliant"
    elif satisfied_count == 0:
        status = "Needs Review"
    else:
        status = "Needs Review"

    return {
        "status": status,
        "satisfied_count": satisfied_count,
        "not_satisfied_count": not_satisfied_count,
        "summary": (
            f"{satisfied_count} guideline item(s) are satisfied and "
            f"{not_satisfied_count} guideline item(s) are non satisfied across Marketing, Legal, and Regulatory guidelines."
        ),
        "sections": sections,
    }
