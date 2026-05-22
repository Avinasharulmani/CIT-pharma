from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import re
import zipfile
import xml.etree.ElementTree as ET

from ..utils import clean_text
from .mlr_embedding_service import embed_text
from .mlr_models import MLRChecklistItem


SECTION_NAME_MAP = {
    "2": "Pre-Submission Checks",
    "3": "Medical Review",
    "4": "Legal Review",
    "5": "Regulatory Review",
    "6": "Final Approval Checks",
}


def read_docx_text(path: Path) -> str:
    try:
        from docx import Document

        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts)
    except Exception:
        pass

    with zipfile.ZipFile(path) as archive:
        xml_data = archive.read("word/document.xml")
    root = ET.fromstring(xml_data)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = []
    for paragraph in root.findall(".//w:p", namespace):
        value = clean_text("".join(node.text or "" for node in paragraph.findall(".//w:t", namespace)))
        if value:
            paragraphs.append(value)
    return "\n".join(paragraphs)


def checklist_version(path: Path) -> str:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
    return f"{path.stem}:{digest}"


def _section_from_heading(line: str) -> str | None:
    match = re.match(r"^Section\s+([2-6])\b.*", line, flags=re.IGNORECASE)
    if match:
        return SECTION_NAME_MAP.get(match.group(1))
    match = re.match(r"^([2-6])\.\s+(.+?)\s*$", line)
    if match and match.group(1) in SECTION_NAME_MAP:
        return SECTION_NAME_MAP[match.group(1)]
    for value in SECTION_NAME_MAP.values():
        if value.lower() in line.lower():
            return value
    return None


def _reviewer_for_section(section_name: str, row_reviewer: str = "") -> str:
    if row_reviewer:
        return row_reviewer
    lowered = section_name.lower()
    if "medical" in lowered:
        return "Medical"
    if "legal" in lowered:
        return "Legal"
    if "regulatory" in lowered:
        return "Regulatory"
    if "final" in lowered:
        return "System / All Reviewers"
    if "pre-submission" in lowered:
        return "Therapy Lead / System"
    return ""


def _looks_like_checklist_row(line: str) -> bool:
    if "|" not in line:
        return False
    parts = [clean_text(part) for part in line.split("|")]
    return len([part for part in parts if part]) >= 3


def _is_header_row(parts: list[str]) -> bool:
    lowered = " ".join(parts).lower()
    return "check item" in lowered and ("mandatory" in lowered or "owner" in lowered or "reviewer" in lowered)


def _parse_row(line: str, section_name: str, row_number: int) -> MLRChecklistItem | None:
    parts = [clean_text(part) for part in line.split("|")]
    parts = [part for part in parts if part]
    if len(parts) < 3:
        return None
    if parts[0].strip("#").lower() in {"item", "check item"} or _is_header_row(parts):
        return None
    item_code = parts[0]
    description = " | ".join(parts[1:-2]) if len(parts) >= 4 else parts[1]
    reviewer = parts[-2] if len(parts) >= 4 else _reviewer_for_section(section_name)
    mandatory = parts[-1]
    if not re.search(r"[A-Za-z]", description):
        return None
    now = datetime.utcnow()
    return MLRChecklistItem(
        item_code=item_code or f"{section_name[:3].upper()}-{row_number}",
        section_name=section_name,
        checklist_description=description,
        reviewer_type=_reviewer_for_section(section_name, reviewer),
        mandatory_status=mandatory,
        checklist_embedding=embed_text(description),
        created_date=now,
        updated_date=now,
    )


def _iter_docx_blocks(path: Path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = clean_text(paragraph.text)
            if text:
                yield "paragraph", text
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = [clean_text(cell.text) for cell in row.cells]
                cells = [cell for cell in cells if cell]
                if cells:
                    yield "row", " | ".join(cells)


def parse_mlr_checklist_docx(path: Path) -> tuple[str, list[MLRChecklistItem]]:
    version = checklist_version(path)
    current_section = ""
    items: list[MLRChecklistItem] = []
    try:
        blocks = list(_iter_docx_blocks(path))
    except Exception:
        blocks = [("paragraph", line) for line in read_docx_text(path).splitlines()]
    for row_number, (_, raw_line) in enumerate(blocks, start=1):
        line = clean_text(raw_line)
        if not line:
            continue
        section_name = _section_from_heading(line)
        if section_name:
            current_section = section_name
            continue
        if not current_section or not _looks_like_checklist_row(line):
            continue
        item = _parse_row(line, current_section, row_number)
        if not item:
            continue
        item.checklist_version = version
        item.source_file_name = path.name
        items.append(item)
    return version, items
