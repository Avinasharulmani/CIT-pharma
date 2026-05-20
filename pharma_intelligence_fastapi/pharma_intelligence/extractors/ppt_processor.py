from pathlib import Path
from typing import List, Optional
import io
import shutil
import subprocess
import tempfile

from ..models import SourceChunk
from ..utils import clean_text
from .image_processor import describe_pil_image_content
from .ocr_utils import run_ocr_on_image
from .text_quality import extraction_unit_metadata, merge_text_blocks


PPT_TEXT_SLIDE_IMAGE_DESCRIPTION_THRESHOLD = 80


def _ppt_slide_image_descriptions(slide, warnings: list[str]) -> tuple[list[str], int]:
    try:
        from PIL import Image
    except Exception:
        return [], 0

    descriptions: list[str] = []
    image_count = 0
    for shape in slide.shapes:
        try:
            image_blob = shape.image.blob
        except Exception:
            continue
        image_count += 1
        try:
            image = Image.open(io.BytesIO(image_blob))
            description = describe_pil_image_content(image, warnings, allow_local_fallback=True)
            if description:
                descriptions.append(f"Image {image_count}: {description}")
        except Exception:
            continue
    return descriptions, image_count


def _ppt_slide_image_count(slide) -> int:
    count = 0
    for shape in slide.shapes:
        try:
            _ = shape.image
        except Exception:
            continue
        count += 1
    return count


def _iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from _iter_shapes(shape.shapes)


def _extract_slide_shape_text(slide) -> str:
    texts = []
    for shape in _iter_shapes(slide.shapes):
        if hasattr(shape, "text") and shape.text:
            texts.append(shape.text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                texts.append(" | ".join(cell.text for cell in row.cells))
    try:
        notes_frame = slide.notes_slide.notes_text_frame
        if notes_frame and notes_frame.text:
            texts.append(notes_frame.text)
    except Exception:
        pass
    return clean_text("\n".join(texts))


def _render_ppt_slides(path: Path, warnings: list[str]) -> dict[int, object]:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        warnings.append("slide rendering skipped because LibreOffice/soffice is not available on PATH")
        return {}
    try:
        import fitz
        from PIL import Image
    except Exception as exc:
        warnings.append(f"slide rendering skipped because PyMuPDF/Pillow is unavailable: {exc}")
        return {}
    try:
        with tempfile.TemporaryDirectory() as folder:
            output_dir = Path(folder)
            command = [
                executable,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(path),
            ]
            result = subprocess.run(command, capture_output=True, text=True, timeout=120)
            if result.returncode != 0:
                warnings.append(f"slide rendering failed: {result.stderr[-400:]}")
                return {}
            pdfs = list(output_dir.glob("*.pdf"))
            if not pdfs:
                warnings.append("slide rendering failed: no PDF output produced")
                return {}
            rendered = {}
            doc = fitz.open(str(pdfs[0]))
            try:
                for index, page in enumerate(doc, start=1):
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    rendered[index] = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
            finally:
                doc.close()
            return rendered
    except Exception as exc:
        warnings.append(f"slide rendering warning: {exc}")
        return {}


def extract_ppt(path: Path, slide_number: Optional[int] = None) -> List[SourceChunk]:
    try:
        from pptx import Presentation
    except Exception as exc:
        raise RuntimeError("python-pptx is required for PPT extraction. Install with: pip install python-pptx") from exc

    prs = Presentation(str(path))
    chunks: List[SourceChunk] = []
    render_warnings: list[str] = []
    rendered_slides = _render_ppt_slides(path, render_warnings)

    for index, slide in enumerate(prs.slides, start=1):
        if slide_number and index != slide_number:
            continue
        warnings: list[str] = list(render_warnings)
        shape_text = _extract_slide_shape_text(slide)
        ocr_text = ""
        ocr_result = None
        if index in rendered_slides:
            try:
                ocr_result = run_ocr_on_image(rendered_slides[index])
                ocr_text = ocr_result.text
                warnings.extend(ocr_result.warnings or [])
            except Exception as exc:
                warnings.append(f"slide OCR warning: {exc}")
        slide_text = merge_text_blocks(shape_text, ocr_text)
        extraction_method = "shape_text_plus_rendered_slide_ocr" if shape_text and ocr_text else ("rendered_slide_ocr" if ocr_text else "shape_text")
        image_count = _ppt_slide_image_count(slide)
        image_descriptions: list[str] = []
        if len(slide_text) < PPT_TEXT_SLIDE_IMAGE_DESCRIPTION_THRESHOLD and (slide_number or len(prs.slides) <= 3):
            image_descriptions, image_count = _ppt_slide_image_descriptions(slide, warnings)
        slide_level = {
            "slide_number": index,
            "shape_text": shape_text,
            "ocr_text": ocr_text,
            "final_text": slide_text,
            "extraction_method": extraction_method,
            "confidence_score": getattr(ocr_result, "confidence_score", None),
            "text_length": len(slide_text),
        }
        chunks.append(
            SourceChunk(
                source_no=index,
                source_type="slide",
                text=slide_text,
                description_of_image_video=clean_text("\n".join(image_descriptions)) or None,
                metadata={
                    **extraction_unit_metadata(
                        source_type="slide",
                        unit_number=index,
                        raw_text=clean_text("\n".join([shape_text, ocr_text])),
                        cleaned_text=slide_text,
                        extraction_method=extraction_method,
                        confidence_score=getattr(ocr_result, "confidence_score", None),
                        warnings=warnings,
                    ),
                    "image_count": image_count,
                    "has_images": image_count > 0,
                    "shape_text": shape_text,
                    "ocr_text": ocr_text,
                    "final_text": slide_text,
                    "slide_level": slide_level,
                    "warnings": warnings,
                },
            )
        )
    return chunks
