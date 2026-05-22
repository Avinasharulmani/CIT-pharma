from pathlib import Path
from typing import Any, Optional
from datetime import datetime
import asyncio
import base64
import csv
import hashlib
import html
import io
import mimetypes
import traceback
import shutil
import subprocess
import textwrap
from uuid import uuid4

from bson import ObjectId
from pymongo import ReturnDocument
from fastapi import BackgroundTasks, Body, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse

from pharma_intelligence.config import (
    ALLOWED_EXTENSIONS,
    BASE_DIR,
    DEFAULT_KEY_MESSAGE_PATH,
    DEFAULT_TARGET_LANGUAGE,
    DEFAULT_TOP_K,
    OUTPUT_DIR,
    SAMPLE_KEY_MESSAGE_PATH,
    SUPPORTED_FILE_TYPES,
    UPLOAD_DIR,
)
from pharma_intelligence.correction.dictionary_loader import fetch_fda_drug_names
from pharma_intelligence.correction.ocr_error_map import seed_default_patterns
from pharma_intelligence.mongo_database import db as mongo_db, init_mongo, review_annotations_collection
from pharma_intelligence.models import BaseIntelligenceResponse
from pharma_intelligence.persistence import save_accuracy_report, save_analysis_response
from pharma_intelligence.permissions import (
    can_add_comment,
    can_delete_comment,
    can_edit_comment,
    can_resolve_comment,
    department_for_role,
    get_current_user_name,
    get_current_user_role,
    is_admin,
    permissions_for_role,
    role_label,
    role_options,
)
from pharma_intelligence.react_ui import REACT_INDEX_HTML
from pharma_intelligence.routes.dictionary import router as dictionary_router
from pharma_intelligence.routes.mlr_routes import router as mlr_router
from pharma_intelligence.services import analyze_file
from pharma_intelligence.services.mlr_review_service import run_mlr_review
from pharma_intelligence.services.mlr_review_engine import run_mlr_review_for_analysis
from pharma_intelligence.extractors.ppt_processor import extract_ppt
from pharma_intelligence.extractors.image_processor import _document_image_segments
from pharma_intelligence.utils import clean_text, detect_file_type, get_file_type_config, save_upload_file
from pharma_intelligence.extractors.generic_processor import (
    extract_document,
    extract_email,
    extract_fallback,
    extract_markup_text,
    extract_spreadsheet,
    extract_zip_package,
)

app = FastAPI(
    title="Pharma Intelligence Engine",
    version="1.0.0",
    description="Reusable FastAPI UI and API for extracting text-based intelligence from pharma commercial materials.",
)
app.include_router(dictionary_router)
app.include_router(mlr_router)

UPLOADED_MATERIALS: dict[str, tuple[Path, str]] = {}
UPLOADED_FILE_HASHES: dict[str, str] = {}
ANALYSIS_JOBS: dict[str, dict[str, Any]] = {}
UPLOAD_IDS_BY_HASH: dict[str, str] = {}
PREVIEW_METADATA_CACHE: dict[tuple[str, str], dict[str, Any]] = {}
PREVIEW_UNITS_CACHE: dict[tuple[str, str, int, int, str], list[dict[str, Any]]] = {}
ANALYSIS_CACHE: dict[tuple[Any, ...], dict[str, Any]] = {}
PREVIEW_DIR = OUTPUT_DIR / "previews"
DIRECT_PREVIEW_EXTENSIONS = {
    ".pdf": "pdf",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".webp": "image",
    ".mp4": "video",
    ".webm": "video",
    ".mov": "video",
    ".m4v": "video",
    ".mp3": "audio",
    ".wav": "audio",
    ".ogg": "audio",
    ".m4a": "audio",
    ".bmp": "image",
    ".avif": "image",
    ".heif": "image",
    ".heic": "image",
    ".tif": "image",
    ".tiff": "image",
    ".svg": "image",
}
CONVERTIBLE_PREVIEW_EXTENSIONS = {".ppt", ".pptx", ".doc", ".docx", ".docm", ".dot", ".dotx", ".dotm", ".rtf"}
WORD_PREVIEW_EXTENSIONS = {".doc", ".docx", ".docm", ".dot", ".dotx", ".dotm", ".rtf"}
TEXT_PREVIEW_EXTENSIONS = {
    ".bash", ".cov", ".csh", ".csv", ".ext", ".htm", ".html", ".inp", ".jsl", ".lst",
    ".odc", ".op", ".param", ".r", ".sas", ".sbml", ".scm", ".sh", ".ssc", ".txt", ".xml",
}


@app.on_event("startup")
async def startup_event():
    init_mongo()
    await seed_default_patterns(mongo_db)
    if mongo_db["pharma_drug_names"].count_documents({"active": True}) == 0:
        await fetch_fda_drug_names(mongo_db)


@app.get("/health")
def health():
    return {
        "status": "ok",
        "database": "mongodb",
        "capabilities": ["supported_file_type_registry", "pdf", "ppt", "documents", "spreadsheets", "images_design", "email", "web_markup_text", "scripts", "zip", "audio", "video", "translation", "key_message_mapping", "visual_annotation"],
        "supported_extensions": sorted(ALLOWED_EXTENSIONS),
    }


def _get_uploaded_material(upload_id: str) -> tuple[Path, str]:
    uploaded = UPLOADED_MATERIALS.get(upload_id)
    if not uploaded:
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        matches = sorted((item for item in UPLOAD_DIR.glob(f"{upload_id}*") if item.is_file()), key=lambda item: item.stat().st_mtime, reverse=True)
        if matches:
            recovered_path = matches[0]
            recovered_name = _display_upload_name(recovered_path)
            UPLOADED_MATERIALS[upload_id] = (recovered_path, recovered_name)
            recovered_hash = _file_hash(recovered_path)
            UPLOADED_FILE_HASHES[upload_id] = recovered_hash
            UPLOAD_IDS_BY_HASH.setdefault(recovered_hash, upload_id)
            return recovered_path, recovered_name
        raise HTTPException(status_code=404, detail="Uploaded file was not found. Please choose the file again.")
    uploaded_path, original_filename = uploaded
    if not uploaded_path.exists():
        raise HTTPException(status_code=404, detail="Uploaded file is no longer available. Please choose the file again.")
    return uploaded_path, original_filename


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _upload_hash(upload_id: str, path: Path | None = None) -> str:
    cached = UPLOADED_FILE_HASHES.get(upload_id)
    if cached:
        return cached
    if path is None:
        path, _ = _get_uploaded_material(upload_id)
    digest = _file_hash(path)
    UPLOADED_FILE_HASHES[upload_id] = digest
    UPLOAD_IDS_BY_HASH.setdefault(digest, upload_id)
    return digest


def _cache_copy(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: _cache_copy(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_cache_copy(item) for item in value]
    return value


def _preview_metadata_for_upload(preview: dict[str, Any], upload_id: str) -> dict[str, Any]:
    item = _cache_copy(preview)
    item["asset_id"] = upload_id
    if "preview_url" in item:
        if item.get("converted"):
            item["preview_url"] = f"/api/materials/{upload_id}/preview-content"
        elif item.get("preview_url"):
            item["preview_url"] = f"/api/materials/{upload_id}/content"
    if "download_url" in item:
        item["download_url"] = f"/api/materials/{upload_id}/download"
    return item


def _json_annotation(document: dict[str, Any]) -> dict[str, Any]:
    item = dict(document)
    item["id"] = str(item.pop("_id"))
    for key in ("created_at", "updated_at"):
        if isinstance(item.get(key), datetime):
            item[key] = item[key].isoformat()
    return item


def _current_user_context(request: Request) -> dict[str, Any]:
    role = get_current_user_role(request.headers)
    user_name = get_current_user_name(request.headers)
    return {
        "user_name": user_name,
        "role": role,
        "role_label": role_label(role),
        "department": department_for_role(role),
        "permissions": permissions_for_role(role),
    }


def _comment_department(document: dict[str, Any]) -> str:
    department = str(document.get("department") or "").strip()
    if department:
        return department
    return department_for_role(document.get("created_by_role"))


@app.get("/api/current-user")
def current_user_api(request: Request):
    context = _current_user_context(request)
    return {**context, "roles": role_options()}


def _office_converter_command() -> list[str] | None:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        return None
    return [executable, "--headless", "--convert-to", "pdf", "--outdir"]


def _preview_file_hash(upload_id: str, source_path: Path) -> str:
    cached = UPLOADED_FILE_HASHES.get(upload_id)
    if cached:
        return cached
    digest = _file_hash(source_path)
    UPLOADED_FILE_HASHES[upload_id] = digest
    UPLOAD_IDS_BY_HASH.setdefault(digest, upload_id)
    return digest


def _converted_pdf_path(upload_id: str, source_path: Path) -> Path:
    file_hash = _preview_file_hash(upload_id, source_path)
    converted_dir = PREVIEW_DIR / "converted"
    converted_dir.mkdir(parents=True, exist_ok=True)
    return converted_dir / f"{file_hash}.pdf"


def _converted_html_path(upload_id: str, source_path: Path) -> Path:
    file_hash = _preview_file_hash(upload_id, source_path)
    converted_dir = PREVIEW_DIR / "converted"
    converted_dir.mkdir(parents=True, exist_ok=True)
    return converted_dir / f"{file_hash}.html"


def _latest_key_message_path() -> Path | None:
    candidates: list[Path] = []
    seen: set[Path] = set()

    for folder in (BASE_DIR, UPLOAD_DIR):
        if not folder.exists():
            continue
        for pattern in ("*.xlsx", "*.xls", "*.csv"):
            for path in folder.glob(pattern):
                if not path.is_file():
                    continue
                if path.resolve() in seen:
                    continue
                seen.add(path.resolve())
                if path.name == SAMPLE_KEY_MESSAGE_PATH.name:
                    continue
                candidates.append(path)

    if not candidates:
        return SAMPLE_KEY_MESSAGE_PATH if SAMPLE_KEY_MESSAGE_PATH.exists() else None

    preferred = sorted(
        candidates,
        key=lambda path: (
            path.resolve() == DEFAULT_KEY_MESSAGE_PATH.resolve() if DEFAULT_KEY_MESSAGE_PATH.exists() else False,
            path.stat().st_mtime,
        ),
        reverse=True,
    )
    return preferred[0]


def _convert_office_to_pdf(upload_id: str, source_path: Path) -> Path | None:
    target_path = _converted_pdf_path(upload_id, source_path)
    if target_path.exists():
        return target_path
    command = _office_converter_command()
    if command:
        temp_dir = PREVIEW_DIR / "conversion_tmp" / uuid4().hex
        temp_dir.mkdir(parents=True, exist_ok=True)
        completed = subprocess.run(
            [*command, str(temp_dir), str(source_path)],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        produced_path = temp_dir / f"{source_path.stem}.pdf"
        if completed.returncode == 0 and produced_path.exists():
            produced_path.replace(target_path)
            shutil.rmtree(temp_dir, ignore_errors=True)
            return target_path
        shutil.rmtree(temp_dir, ignore_errors=True)
    if source_path.suffix.lower() in WORD_PREVIEW_EXTENSIONS:
        return _convert_word_to_pdf(upload_id, source_path)
    return None


def _convert_word_to_pdf(upload_id: str, source_path: Path) -> Path | None:
    target_path = _converted_pdf_path(upload_id, source_path)
    if target_path.exists():
        return target_path
    source = _powershell_quote(str(source_path.resolve()))
    target = _powershell_quote(str(target_path.resolve()))
    script = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
$document = $null
try {{
  $document = $word.Documents.Open({source}, $false, $true, $false)
  $document.ExportAsFixedFormat({target}, 17)
}} finally {{
  if ($document -ne $null) {{
    $document.Close(0)
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($document) | Out-Null
  }}
  $word.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
}}
"""
    completed = subprocess.run(
        ["powershell", "-NoProfile", "-STA", "-ExecutionPolicy", "Bypass", "-Command", script],
        capture_output=True,
        text=True,
        timeout=90,
        check=False,
    )
    if completed.returncode == 0 and target_path.exists():
        return target_path
    return None


def _extract_preview_text(source_path: Path, original_filename: str, max_chars: Optional[int] = None) -> tuple[str, list[dict[str, Any]], list[str]]:
    extension = source_path.suffix.lower()
    file_type = detect_file_type(original_filename)
    warnings: list[str] = []
    try:
        if file_type == "document":
            chunks = extract_document(source_path)
        elif file_type == "ppt":
            chunks = extract_ppt(source_path)
        elif file_type == "spreadsheet":
            chunks = extract_spreadsheet(source_path)
        elif file_type == "email":
            chunks = extract_email(source_path)
        elif file_type in {"text", "script"} or extension == ".svg":
            chunks = extract_markup_text(source_path)
        elif file_type == "archive":
            chunks = extract_zip_package(source_path)
        else:
            chunks = extract_fallback(source_path)
    except Exception as exc:
        chunks = extract_fallback(source_path)
        warnings.append(f"Preview extraction warning: {exc}")
    for chunk in chunks:
        warnings.extend(chunk.metadata.get("warnings", []) or [])
    def preview_text(value: str) -> str:
        text_value = clean_text(value)
        return text_value[:max_chars] if max_chars is not None else text_value

    units = [
        {
            "number": chunk.source_no,
            "label": f"{chunk.source_type.title()} {chunk.source_no}" if chunk.source_no else chunk.source_type.title(),
            "title": str(chunk.metadata.get("title") or chunk.metadata.get("sheet_name") or chunk.metadata.get("member_name") or chunk.source_type.title()),
            "text": preview_text(chunk.text),
            "metadata": chunk.metadata,
        }
        for chunk in chunks
    ]
    _add_textual_preview_thumbnails(units, file_type)
    text = preview_text("\n\n".join(unit["text"] for unit in units if unit["text"]))
    return text, units, warnings


def _write_html_preview(upload_id: str, source_path: Path, original_filename: str, text: str, units: list[dict[str, Any]]) -> Path | None:
    if not text and not units:
        return None
    target_path = _converted_html_path(upload_id, source_path)
    sections: list[str] = []
    for index, unit in enumerate(units, start=1):
        unit_text = clean_text(str(unit.get("text") or unit.get("raw_text") or ""))
        if not unit_text:
            continue
        label = str(unit.get("label") or unit.get("title") or f"Section {index}")
        sections.append(
            "<section class=\"preview-section\">"
            f"<h2>{html.escape(label)}</h2>"
            f"<pre>{html.escape(unit_text)}</pre>"
            "</section>"
        )
    if not sections and text:
        sections.append(f"<section class=\"preview-section\"><pre>{html.escape(clean_text(text))}</pre></section>")
    document = f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>{html.escape(original_filename)}</title>
  <style>
    body {{ margin: 0; padding: 28px; background: #f7f9fc; color: #17233d; font-family: Arial, sans-serif; }}
    .preview-shell {{ max-width: 980px; margin: 0 auto; }}
    h1 {{ margin: 0 0 18px; font-size: 24px; }}
    .preview-section {{ margin: 0 0 18px; padding: 18px; background: #fff; border: 1px solid #dfe6f1; border-radius: 8px; }}
    h2 {{ margin: 0 0 12px; font-size: 16px; color: #1f4ed8; }}
    pre {{ margin: 0; white-space: pre-wrap; word-break: break-word; font: 15px/1.55 Arial, sans-serif; user-select: text; }}
  </style>
</head>
<body>
  <main class="preview-shell">
    <h1>{html.escape(original_filename)}</h1>
    {''.join(sections)}
  </main>
</body>
</html>
"""
    target_path.write_text(document, encoding="utf-8")
    return target_path


def _convert_material_to_html_preview(upload_id: str, source_path: Path, original_filename: str) -> tuple[Path | None, str, list[dict[str, Any]], list[str]]:
    text, units, warnings = _extract_preview_text(source_path, original_filename)
    html_path = _write_html_preview(upload_id, source_path, original_filename, text, units)
    return html_path, text, units, warnings


def _add_textual_preview_thumbnails(units: list[dict[str, Any]], file_type: str) -> None:
    for unit in units:
        text = str(unit.get("text") or "")
        if not text:
            continue
        source_type = str(unit.get("metadata", {}).get("source_type") or unit.get("label") or "").lower()
        is_table = file_type == "spreadsheet" or "sheet" in source_type or "table" in source_type
        if is_table:
            unit["table"] = _preview_table_rows(text, max_rows=250)
        thumbnail = _table_preview_thumbnail(text, str(unit.get("title") or unit.get("label") or "Sheet")) if is_table else _text_preview_thumbnail(text, str(unit.get("title") or unit.get("label") or "Text"))
        if thumbnail:
            unit["thumbnail"] = thumbnail


def _text_preview_thumbnail(text: str, title: str) -> str:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return ""
    width, height = 900, 1180
    image = Image.new("RGB", (width, height), "#ffffff")
    draw = ImageDraw.Draw(image)
    font_title = _preview_font(28, bold=True)
    font_body = _preview_font(22)
    draw.rectangle((0, 0, width - 1, height - 1), outline="#dbe3ef", width=2)
    draw.text((56, 48), title[:80], fill="#17346d", font=font_title)
    y = 108
    for paragraph in str(text).replace("\r", "\n").splitlines():
        line = clean_text(paragraph)
        if not line:
            y += 16
            continue
        for wrapped in textwrap.wrap(line, width=72, break_long_words=False, replace_whitespace=False) or [""]:
            if y > height - 70:
                break
            draw.text((56, y), wrapped, fill="#17233d", font=font_body)
            y += 31
        if y > height - 70:
            break
    return _image_to_data_url(image)


def _table_preview_thumbnail(text: str, title: str) -> str:
    try:
        from PIL import Image, ImageDraw
    except Exception:
        return ""
    rows = _preview_table_rows(text, max_rows=30)
    if not rows:
        return _text_preview_thumbnail(text, title)
    width, height = 1100, 760
    image = Image.new("RGB", (width, height), "#ffffff")
    draw = ImageDraw.Draw(image)
    font_title = _preview_font(26, bold=True)
    font_header = _preview_font(18, bold=True)
    font_cell = _preview_font(17)
    draw.rectangle((0, 0, width - 1, height - 1), outline="#dbe3ef", width=2)
    draw.text((34, 26), title[:90], fill="#17346d", font=font_title)
    max_cols = min(max(len(row) for row in rows), 6)
    shown_rows = rows[:14]
    table_left, table_top = 34, 82
    table_width = width - 68
    col_width = max(100, table_width // max_cols)
    row_height = 44
    for row_index, row in enumerate(shown_rows):
        y = table_top + row_index * row_height
        if y + row_height > height - 28:
            break
        is_header = row_index == 0
        fill = "#eef4ff" if is_header else ("#ffffff" if row_index % 2 else "#f8fbff")
        draw.rectangle((table_left, y, table_left + table_width, y + row_height), fill=fill, outline="#dbe3ef")
        for col_index in range(max_cols):
            x = table_left + col_index * col_width
            draw.line((x, y, x, y + row_height), fill="#dbe3ef")
            value = clean_text(row[col_index] if col_index < len(row) else "")
            draw.text((x + 10, y + 12), _fit_preview_text(value, 18), fill="#17233d", font=font_header if is_header else font_cell)
        draw.line((table_left + max_cols * col_width, y, table_left + max_cols * col_width, y + row_height), fill="#dbe3ef")
    return _image_to_data_url(image)


def _preview_table_rows(text: str, max_rows: int = 250) -> list[list[str]]:
    sample = str(text or "").strip()
    if not sample:
        return []
    try:
        reader = csv.reader(io.StringIO(sample))
        rows = [[clean_text(cell) for cell in row] for row in reader if any(clean_text(cell) for cell in row)]
    except Exception:
        rows = []
    if len(rows) <= 1:
        rows = [[clean_text(cell) for cell in line.split("|")] for line in sample.splitlines() if clean_text(line)]
    return [row[:50] for row in rows[:max_rows]]


def _preview_font(size: int, bold: bool = False):
    try:
        from PIL import ImageFont
        font_name = "arialbd.ttf" if bold else "arial.ttf"
        return ImageFont.truetype(font_name, size)
    except Exception:
        from PIL import ImageFont
        return ImageFont.load_default()


def _fit_preview_text(value: str, max_chars: int) -> str:
    value = clean_text(value)
    if len(value) <= max_chars:
        return value
    return value[: max(1, max_chars - 1)].rstrip() + "..."


def _image_to_data_url(image) -> str:
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return f"data:image/png;base64,{base64.b64encode(buffer.getvalue()).decode('ascii')}"


def _fallback_preview_card(upload_id: str, source_path: Path, original_filename: str, message: str | None = None) -> dict[str, Any]:
    file_info = get_file_type_config(original_filename)
    extension = file_info.get("extension") or source_path.suffix.lower()
    return {
        "asset_id": upload_id,
        "file_name": original_filename,
        "extension": extension,
        "category": file_info.get("category"),
        "kind": "fallback",
        "preview_type": "fallback",
        "preview_url": "",
        "download_url": f"/api/materials/{upload_id}/download",
        "preview_available": False,
        "converted": False,
        "message": message or "Preview is not available for this file type. The system will still attempt content extraction and mapping.",
        "analysis_status": "Ready for analysis",
        "mapping_status": "Ready for mapping",
    }


def _powershell_quote(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _powerpoint_preview_units(upload_id: str, path: Path, start: int = 1, count: int | None = None, quality: str = "thumb") -> list[dict[str, Any]]:
    file_hash = _preview_file_hash(upload_id, path)
    export_dir = PREVIEW_DIR / "ppt_slides" / f"{file_hash}_{quality}"
    export_dir.mkdir(parents=True, exist_ok=True)
    safe_start = max(1, int(start or 1))
    safe_count = None if count is None else max(1, int(count or 1))
    width = 960 if quality != "full" else 1920
    height = 540 if quality != "full" else 1080
    requested_numbers = list(range(safe_start, safe_start + (safe_count or 1)))
    def exported_slide_number(item: Path) -> int | None:
        try:
            return int(item.stem.split("_", 1)[1])
        except Exception:
            return None

    existing = sorted(path for path in export_dir.glob("slide_*.png") if exported_slide_number(path) in requested_numbers)
    if len(existing) < len(requested_numbers):
        source = _powershell_quote(str(path.resolve()))
        target = _powershell_quote(str(export_dir.resolve()))
        slide_numbers = ",".join(str(number) for number in requested_numbers)
        script = f"""
$ErrorActionPreference = 'Stop'
$ppt = New-Object -ComObject PowerPoint.Application
$presentation = $ppt.Presentations.Open({source}, $true, $false, $false)
try {{
  $slideNumbers = @({slide_numbers})
  foreach ($slideNumber in $slideNumbers) {{
    if ($slideNumber -ge 1 -and $slideNumber -le $presentation.Slides.Count) {{
      $slide = $presentation.Slides.Item($slideNumber)
      $slide.Export((Join-Path {target} ('slide_' + $slide.SlideIndex + '.png')), 'PNG', {width}, {height})
    }}
  }}
}} finally {{
  $presentation.Close()
  $ppt.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($presentation) | Out-Null
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($ppt) | Out-Null
}}
"""
        completed = subprocess.run(
            ["powershell", "-NoProfile", "-STA", "-ExecutionPolicy", "Bypass", "-Command", script],
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
        )
        if completed.returncode != 0:
            return []
        existing = sorted(path for path in export_dir.glob("slide_*.png") if exported_slide_number(path) in requested_numbers)
    if not existing:
        return []

    slide_text_by_number: dict[int, str] = {}
    try:
        from pptx import Presentation
        presentation = Presentation(str(path))
        for slide_index, slide in enumerate(presentation.slides, start=1):
            slide_text_by_number[slide_index] = _ppt_slide_text(slide)
    except Exception:
        slide_text_by_number = {}

    units: list[dict[str, Any]] = []
    for image_path in existing:
        try:
            slide_number = int(image_path.stem.split("_", 1)[1])
        except Exception:
            continue
        image_data = image_path.read_bytes()
        image_width = width
        image_height = height
        try:
            from PIL import Image
            with Image.open(image_path) as image:
                image_width, image_height = image.size
                image_text = _preview_ocr_text(image)
        except Exception:
            image_text = ""
            pass
        units.append(
            {
                "number": slide_number,
                "label": f"Slide {slide_number}",
                "title": "",
                "text": slide_text_by_number.get(slide_number, "") or image_text,
                "width": image_width,
                "height": image_height,
                "quality": quality,
                "thumbnail": f"data:image/png;base64,{base64.b64encode(image_data).decode('ascii')}",
            }
        )
    return sorted(units, key=lambda unit: unit["number"])


def _preview_metadata(upload_id: str, source_path: Path, original_filename: str) -> dict[str, Any]:
    extension = source_path.suffix.lower()
    file_info = get_file_type_config(original_filename)
    download_url = f"/api/materials/{upload_id}/download"
    if extension in DIRECT_PREVIEW_EXTENSIONS:
        return {
            "asset_id": upload_id,
            "file_name": original_filename,
            "extension": extension,
            "category": file_info.get("category"),
            "kind": DIRECT_PREVIEW_EXTENSIONS[extension],
            "preview_type": file_info.get("preview_type"),
            "preview_url": f"/api/materials/{upload_id}/content",
            "download_url": download_url,
            "preview_available": True,
            "converted": False,
            "analysis_status": "Ready for analysis",
            "mapping_status": "Ready for mapping",
        }
    if extension in CONVERTIBLE_PREVIEW_EXTENSIONS:
        converted_path = _convert_office_to_pdf(upload_id, source_path)
        if converted_path and converted_path.exists():
            units = _pdf_preview_units(converted_path, label_prefix="Slide" if file_info.get("file_type") == "ppt" else "Page")
            return {
                "asset_id": upload_id,
                "file_name": original_filename,
                "extension": extension,
                "category": file_info.get("category"),
                "kind": file_info.get("file_type") or "document",
                "preview_type": file_info.get("preview_type") or file_info.get("file_type") or "document",
                "preview_url": "",
                "download_url": download_url,
                "preview_available": True,
                "converted": False,
                "units": units[:20],
                "message": "Visual preview generated from the uploaded file.",
                "analysis_status": "Ready for analysis",
                "mapping_status": "Ready for mapping",
            }
        if file_info.get("preview_type") == "text":
            text, units, warnings = _extract_preview_text(source_path, original_filename)
            return {
                "asset_id": upload_id,
                "file_name": original_filename,
                "extension": extension,
                "category": file_info.get("category"),
                "kind": "text",
                "preview_type": "text",
                "preview_url": "",
                "download_url": download_url,
                "preview_available": bool(text),
                "converted": False,
                "text_preview": text,
                "units": units,
                "warnings": warnings,
                "message": "Text preview generated from extracted content." if text else "Preview conversion is unavailable and no readable preview text was found.",
                "analysis_status": "Ready for analysis",
                "mapping_status": "Ready for mapping",
            }
        return {
            **_fallback_preview_card(upload_id, source_path, original_filename, "Preview conversion is unavailable on this server. Download is still available."),
        }
    if extension in TEXT_PREVIEW_EXTENSIONS or file_info.get("preview_type") in {"text", "code", "table", "email", "package", "html"}:
        text, units, warnings = _extract_preview_text(source_path, original_filename)
        unit_limit = None if file_info.get("preview_type") in {"text", "code", "html"} or extension in TEXT_PREVIEW_EXTENSIONS else 20
        return {
            "asset_id": upload_id,
            "file_name": original_filename,
            "extension": extension,
            "category": file_info.get("category"),
            "kind": file_info.get("preview_type") or "text",
            "preview_type": file_info.get("preview_type") or "text",
            "preview_url": "",
            "download_url": download_url,
            "preview_available": bool(text or units),
            "converted": False,
            "text_preview": text,
            "units": units if unit_limit is None else units[:unit_limit],
            "warnings": warnings,
            "message": "Preview generated from extracted content." if text or units else "No readable preview content was found.",
            "analysis_status": "Ready for analysis",
            "mapping_status": "Ready for mapping",
        }
    if file_info.get("preview_type") == "fallback":
        text, units, warnings = _extract_preview_text(source_path, original_filename)
        if text:
            return {
                "asset_id": upload_id,
                "file_name": original_filename,
                "extension": extension,
                "category": file_info.get("category"),
                "kind": "text",
                "preview_type": "text",
                "preview_url": "",
                "download_url": download_url,
                "preview_available": True,
                "converted": False,
                "text_preview": text,
                "units": units,
                "warnings": warnings,
                "message": "Direct visual preview is not available, but readable text was extracted for preview.",
                "analysis_status": "Ready for analysis",
                "mapping_status": "Ready for mapping",
            }
    return {
        "asset_id": upload_id,
        "file_name": original_filename,
        "extension": extension,
        "category": file_info.get("category"),
        "kind": "fallback",
        "preview_type": "fallback",
        "preview_url": "",
        "download_url": download_url,
        "preview_available": False,
        "converted": False,
        "message": "This file type cannot be previewed directly. The system will still attempt extraction and mapping.",
        "analysis_status": "Ready for analysis",
        "mapping_status": "Ready for mapping",
    }


def _upload_preview_metadata(upload_id: str, source_path: Path, original_filename: str) -> dict[str, Any]:
    extension = source_path.suffix.lower()
    file_info = get_file_type_config(original_filename)
    direct_kind = DIRECT_PREVIEW_EXTENSIONS.get(extension)
    if extension in CONVERTIBLE_PREVIEW_EXTENSIONS:
        return {
            "asset_id": upload_id,
            "file_name": original_filename,
            "extension": extension,
            "category": file_info.get("category"),
            "kind": file_info.get("file_type") or "document",
            "file_type": file_info.get("file_type") or "document",
            "preview_type": file_info.get("preview_type") or file_info.get("file_type") or "document",
            "preview_url": "",
            "download_url": f"/api/materials/{upload_id}/download",
            "preview_available": True,
            "converted": False,
            "message": "Visual preview will be rendered on demand.",
            "analysis_status": "Ready for analysis",
            "mapping_status": "Ready for mapping",
        }
    return {
        "asset_id": upload_id,
        "file_name": original_filename,
        "extension": extension,
        "category": file_info.get("category"),
        "kind": direct_kind or file_info.get("file_type") or "document",
        "preview_type": file_info.get("preview_type"),
        "preview_url": f"/api/materials/{upload_id}/content" if direct_kind else "",
        "download_url": f"/api/materials/{upload_id}/download",
        "preview_available": bool(direct_kind or file_info.get("can_preview")),
        "converted": False,
        "message": "Preview will be rendered on demand.",
        "analysis_status": "Ready for analysis",
        "mapping_status": "Ready for mapping",
    }


def _pdf_preview_units(path: Path, label_prefix: str = "Page", start: int = 1, count: int | None = None, quality: str = "thumb") -> list[dict[str, Any]]:
    try:
        import fitz
    except Exception as exc:
        raise RuntimeError(f"PyMuPDF is required for PDF preview: {exc}") from exc

    units: list[dict[str, Any]] = []
    safe_start = max(1, int(start or 1))
    safe_count = None if count is None else max(1, int(count or 1))
    scale = 0.55 if quality != "full" else 2.0
    doc = fitz.open(str(path))
    try:
        end = len(doc) if safe_count is None else min(len(doc), safe_start + safe_count - 1)
        for index in range(safe_start, end + 1):
            page = doc[index - 1]
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            image_data = pixmap.tobytes("png")
            units.append(
                {
                    "number": index,
                    "label": f"{label_prefix} {index}",
                    "text": clean_text(page.get_text("text") or ""),
                    "width": pixmap.width,
                    "height": pixmap.height,
                    "quality": quality,
                    "text_boxes": _pdf_text_boxes(page),
                    "thumbnail": f"data:image/png;base64,{base64.b64encode(image_data).decode('ascii')}",
                }
            )
    finally:
        doc.close()
    return units


def _pdf_text_boxes(page: Any) -> list[dict[str, Any]]:
    page_rect = page.rect
    page_width = float(page_rect.width or 1)
    page_height = float(page_rect.height or 1)
    boxes: list[dict[str, Any]] = []
    try:
        page_dict = page.get_text("dict")
    except Exception:
        return boxes
    for block in page_dict.get("blocks", []) or []:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []) or []:
            spans = line.get("spans", []) or []
            text = clean_text(" ".join(str(span.get("text") or "") for span in spans))
            if not text:
                continue
            bbox = line.get("bbox") or (spans[0].get("bbox") if spans else None)
            if not bbox or len(bbox) < 4:
                continue
            x0, y0, x1, y1 = [float(value) for value in bbox[:4]]
            width = max(0.1, ((x1 - x0) / page_width) * 100)
            height = max(0.1, ((y1 - y0) / page_height) * 100)
            if width <= 0 or height <= 0:
                continue
            font_size = max((float(spans[0].get("size") or 10) / page_width) * 100, 0.4) if spans else height
            boxes.append(
                {
                    "text": text,
                    "x": max(0.0, min(100.0, (x0 / page_width) * 100)),
                    "y": max(0.0, min(100.0, (y0 / page_height) * 100)),
                    "width": min(100.0, width),
                    "height": min(100.0, height),
                    "font_size": min(8.0, font_size),
                }
            )
    return boxes


def _video_preview_units(path: Path) -> list[dict[str, Any]]:
    try:
        import cv2
    except Exception:
        return []

    capture = cv2.VideoCapture(str(path))
    if not capture.isOpened():
        return []
    units: list[dict[str, Any]] = []
    try:
        fps = float(capture.get(cv2.CAP_PROP_FPS) or 0)
        frame_count = float(capture.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
        duration = frame_count / fps if fps > 0 and frame_count > 0 else 0
        total_seconds = max(1, int(duration))
        max_frames = 120
        if total_seconds <= max_frames:
            seconds = list(range(0, total_seconds + 1))
        else:
            step = total_seconds / max_frames
            seconds = sorted({int(round(index * step)) for index in range(max_frames + 1)})
        for second in seconds:
            capture.set(cv2.CAP_PROP_POS_MSEC, max(0, second) * 1000)
            ok, frame = capture.read()
            if (not ok or frame is None) and second == 0:
                capture.set(cv2.CAP_PROP_POS_MSEC, 250)
                ok, frame = capture.read()
            if not ok or frame is None:
                continue
            ok, buffer = cv2.imencode(".jpg", frame, [int(cv2.IMWRITE_JPEG_QUALITY), 78])
            if not ok:
                continue
            units.append(
                {
                    "number": len(units) + 1,
                    "label": _format_preview_seconds(second),
                    "second": second,
                    "thumbnail": f"data:image/jpeg;base64,{base64.b64encode(buffer.tobytes()).decode('ascii')}",
                }
            )
    finally:
        capture.release()
    return units


def _preview_ocr_text(image: Any) -> str:
    try:
        from pharma_intelligence.extractors.image_processor import _ocr_pil_image
        text, _ = _ocr_pil_image(image)
        return clean_text(text or "")
    except Exception:
        return ""


def _image_preview_units(path: Path) -> list[dict[str, Any]]:
    try:
        from PIL import Image
    except Exception:
        return []
    try:
        image = Image.open(str(path))
        segments = _document_image_segments(image)
    except Exception:
        return []
    units: list[dict[str, Any]] = []
    for number, segment in segments:
        try:
            preview = segment.convert("RGB")
            preview.thumbnail((900, 1200))
            buffer = io.BytesIO()
            preview.save(buffer, format="PNG")
            units.append(
                {
                    "number": number,
                    "label": f"Page {number}" if len(segments) > 1 else "Image",
                    "title": f"Page {number}" if len(segments) > 1 else "Image",
                    "text": _preview_ocr_text(segment),
                    "width": segment.size[0],
                    "height": segment.size[1],
                    "thumbnail": f"data:image/png;base64,{base64.b64encode(buffer.getvalue()).decode('ascii')}",
                }
            )
        except Exception:
            continue
    return units


def _format_preview_seconds(seconds: int) -> str:
    minutes = max(0, int(seconds)) // 60
    remaining = max(0, int(seconds)) % 60
    return f"{minutes}:{remaining:02d}"


def _wrap_svg_text(text: str, max_chars: int, max_lines: int = 5) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) <= max_chars:
            current = candidate
            continue
        if current:
            lines.append(current)
        current = word
        if len(lines) >= max_lines:
            break
    if current and len(lines) < max_lines:
        lines.append(current)
    return lines


def _ppt_slide_thumbnail(slide, slide_width: int, slide_height: int) -> str:
    width = 1440
    height = max(1, round(width * (slide_height / slide_width)))
    scale_x = width / slide_width
    scale_y = height / slide_height
    elements = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#ffffff"/>',
    ]
    for shape in slide.shapes:
        try:
            left = float(shape.left)
            top = float(shape.top)
            shape_width = max(1, float(shape.width))
            shape_height = max(1, float(shape.height))
        except Exception:
            continue
        x = left * scale_x
        y = top * scale_y
        w = max(1, shape_width * scale_x)
        h = max(1, shape_height * scale_y)

        image = getattr(shape, "image", None)
        if image is not None:
            image_data = base64.b64encode(image.blob).decode("ascii")
            image_type = image.content_type or "image/png"
            elements.append(
                f'<image x="{x:.2f}" y="{y:.2f}" width="{w:.2f}" height="{h:.2f}" '
                f'preserveAspectRatio="none" href="data:{html.escape(image_type)};base64,{image_data}"/>'
            )
            continue

        text = clean_text(getattr(shape, "text", "") or "")
        if not text:
            continue
        elements.append(
            f'<rect x="{x:.2f}" y="{y:.2f}" width="{w:.2f}" height="{h:.2f}" '
            'rx="6" fill="rgba(248,250,252,0.82)" stroke="#d7deea" stroke-width="1"/>'
        )
        font_size = max(12, min(26, h / 5.5))
        text_x = x + min(18, w * 0.06)
        text_y = y + min(28, h * 0.18)
        max_chars = max(12, int(w / (font_size * 0.55)))
        elements.append(
            f'<text x="{text_x:.2f}" y="{text_y:.2f}" fill="#172033" '
            f'font-family="Arial, sans-serif" font-size="{font_size:.2f}" font-weight="650">'
        )
        for line_index, line in enumerate(_wrap_svg_text(text, max_chars=max_chars, max_lines=max(1, int(h / (font_size * 1.3))))):
            dy = 0 if line_index == 0 else font_size * 1.25
            elements.append(f'<tspan x="{text_x:.2f}" dy="{dy:.2f}">{html.escape(line)}</tspan>')
        elements.append("</text>")
    elements.append("</svg>")
    svg = "".join(elements)
    return f"data:image/svg+xml;base64,{base64.b64encode(svg.encode('utf-8')).decode('ascii')}"


def _ppt_slide_text(slide: Any) -> str:
    parts: list[str] = []
    for shape in getattr(slide, "shapes", []) or []:
        text = getattr(shape, "text", "")
        if text:
            parts.append(str(text))
        if getattr(shape, "has_table", False):
            try:
                for row in shape.table.rows:
                    row_text = " ".join(str(cell.text or "").strip() for cell in row.cells if str(cell.text or "").strip())
                    if row_text:
                        parts.append(row_text)
            except Exception:
                pass
    return clean_text("\n".join(parts))


def _ppt_preview_units(path: Path, start: int = 1, count: int | None = None) -> list[dict[str, Any]]:
    try:
        from pptx import Presentation
    except Exception as exc:
        raise RuntimeError(f"python-pptx is required for PPT preview: {exc}") from exc

    prs = Presentation(str(path))
    units: list[dict[str, Any]] = []
    safe_start = max(1, int(start or 1))
    safe_count = None if count is None else max(1, int(count or 1))
    end = len(prs.slides) if safe_count is None else min(len(prs.slides), safe_start + safe_count - 1)
    for index in range(safe_start, end + 1):
        slide = prs.slides[index - 1]
        units.append(
            {
                "number": index,
                "label": f"Slide {index}",
                "title": "",
                "text": "",
                "width": 1440,
                "height": max(1, round(1440 * (prs.slide_height / prs.slide_width))),
                "thumbnail": _ppt_slide_thumbnail(slide, prs.slide_width, prs.slide_height),
            }
        )
    return units


def _layout(content: str) -> str:
    return f"""
    <!DOCTYPE html>
    <html>
    <head>
      <title>Pharma Intelligence Engine</title>
      <style>
        body {{ font-family: Arial, sans-serif; background: #f6f8fb; margin: 0; color: #1f2937; }}
        .container {{ max-width: 1180px; margin: 28px auto; background: #fff; border-radius: 16px; padding: 28px; box-shadow: 0 8px 24px rgba(15,23,42,.08); }}
        h1 {{ margin: 0 0 8px 0; }}
        .subtitle {{ color: #4b5563; line-height: 1.5; margin-bottom: 22px; }}
        .notice {{ background: #eef6ff; border-left: 4px solid #2563eb; padding: 12px 14px; border-radius: 10px; margin: 16px 0; }}        .section { background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 8px; padding: 16px; margin: 16px 0; }
        .section h3 { margin: 0 0 12px 0; color: #1e293b; }
        .content-box { background: #fff; border: 1px solid #cbd5e1; border-radius: 6px; padding: 12px; max-height: 300px; overflow-y: auto; white-space: pre-wrap; font-size: 14px; line-height: 1.5; }        .warning {{ background: #fffbeb; color: #92400e; border-left: 4px solid #f59e0b; padding: 12px 14px; border-radius: 10px; margin: 16px 0; }}
        .grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 16px; }}
        label {{ display: block; font-weight: 700; margin: 12px 0 6px; }}
        input, select {{ width: 100%; box-sizing: border-box; padding: 10px; border: 1px solid #d1d5db; border-radius: 10px; }}
        input[type='checkbox'] {{ width: auto; }}
        button {{ margin-top: 18px; padding: 12px 18px; border: 0; border-radius: 10px; background: #2563eb; color: #fff; font-weight: 700; cursor: pointer; }}
        table {{ width: 100%; border-collapse: collapse; margin-top: 18px; font-size: 14px; }}
        th, td {{ border: 1px solid #e5e7eb; padding: 10px; vertical-align: top; }}
        th {{ background: #f3f4f6; text-align: left; }}
        .pill {{ display: inline-block; background: #ecfdf5; color: #065f46; border-radius: 999px; padding: 3px 8px; margin: 3px 3px 3px 0; }}
        pre {{ white-space: pre-wrap; background: #f9fafb; padding: 16px; border-radius: 10px; border: 1px solid #e5e7eb; }}
        .error {{ background: #fef2f2; color: #991b1b; border-left: 4px solid #dc2626; padding: 12px; border-radius: 10px; }}
      </style>
    </head>
    <body><div class="container">{content}</div></body>
    </html>
    """


@app.get("/", response_class=HTMLResponse)
def index():
    return HTMLResponse(REACT_INDEX_HTML, headers={"Cache-Control": "no-store, max-age=0", "Pragma": "no-cache"})

    content = """
    <h1>💊 Pharma Intelligence Engine</h1>
    <p class="subtitle">
      Extract text-based intelligence from Pharma commercial materials (PPT, PDF, Video, Image, Audio).
      Upload a key-message Excel/CSV file, or leave it empty to use the project default.
    </p>
    <div class="notice">
      <b>Workflow:</b> upload → auto-detect file type → intelligence extraction → optional translation → key-message mapping → result.
    </div>
    <form action="/analyze" method="post" enctype="multipart/form-data" id="uploadForm">
      <div class="grid">
        <div style="grid-column: span 2;">
          <label>Upload Pharma Commercial Material</label>
          <input type="file" name="file" id="fileInput" accept=".eps,.ai,.indd,.zip,.psd,.avif,.bmp,.cov,.ext,.inp,.lst,.op,.param,.r,.sas,.scm,.jsl,.ssc,.sbml,.eml,.gif,.heif,.heic,.html,.htm,.jpg,.jpeg,.xls,.xlsx,.odc,.msg,.ppt,.pptx,.doc,.docx,.docm,.dotm,.dot,.dotx,.pdf,.png,.raw,.dng,.crw,.cr2,.raf,.3fr,.dcr,.kdc,.nef,.nrw,.orf,.rw2,.pef,.arw,.srf,.sr2,.rtf,.bash,.csh,.sh,.svg,.tif,.tiff,.txt,.csv,.webp,.xml,.hwp,.mp3,.wav,.m4a,.flac,.ogg,.mp4,.avi,.mov,.mkv" required>
        </div>
        <div style="grid-column: span 2;">
          <label>Key Message File (optional)</label>
          <input type="file" name="key_message_file" accept=".xlsx,.xls,.csv">
        </div>
      </div>
      
      <div class="grid">
        <div id="page_slide_container" style="display: none;">
          <label>Page / Slide Number (optional for PDF/PPT)</label>
          <input type="number" name="page_slide_number" min="1" placeholder="Example: 3">
        </div>
        <div>
           <label>Target Language</label>
           <input type="text" name="target_language" value="en" placeholder="en, ta, hi, fr...">
        </div>
      </div>
      <div class="grid">
        <div>
          <label>Top Matches (0 = all)</label>
          <input type="number" name="top_k" min="0" max="100" value="0">
        </div>
        <div>
          <label>Mapping Enabled</label>
          <input type="checkbox" name="mapping_enabled" checked> Map extracted content with key messages
        </div>
      </div>
      
      <button type="submit" id="submitBtn">Analyze Document</button>
    </form>

    <script>
      const fileInput = document.getElementById('fileInput');
      const pageSlideContainer = document.getElementById('page_slide_container');

      fileInput.addEventListener('change', function() {
        if (this.files && this.files[0]) {
          const fileName = this.files[0].name.toLowerCase();
          if (fileName.endsWith('.pdf') || fileName.endsWith('.pptx')) {
            pageSlideContainer.style.display = 'block';
          } else {
            pageSlideContainer.style.display = 'none';
          }
        } else {
          pageSlideContainer.style.display = 'none';
        }
      });
    </script>
    """
    return _layout(content)


def _render_result(response: BaseIntelligenceResponse) -> HTMLResponse:
    keywords = "".join(f"<span class='pill'>{html.escape(k)}</span>" for k in response.extracted_keywords)
    description = response.description_of_image_video or "N/A"
    warnings = response.metadata.get("warnings") or []
    warning_html = "".join(f"<div>{html.escape(str(warning))}</div>" for warning in warnings)
    warning_block = f"<div class='warning'>{warning_html}</div>" if warning_html else ""
    
    # Build pre-table content based on file type
    pre_table_content = ""
    
    if response.file_type in {"video", "audio"}:
        # Show transcript and frame descriptions before table
        transcript_text = response.full_transcript or response.transcript
        transcript_section = ""
        if transcript_text:
            transcript_section = f"""
        <div class='section'>
            <h3>🎤 Transcript / Audio Content</h3>
            <div class='content-box'>{html.escape(transcript_text)}</div>
        </div>"""
        
        frame_section = ""
        if response.description_of_image_video:
            frame_section = f"""
        <div class='section'>
            <h3>🎬 Visual Description</h3>
            <div class='content-box'>{html.escape(response.description_of_image_video)}</div>
        </div>"""
        
        pre_table_content = transcript_section + frame_section
    
    elif response.file_type in {"pdf", "ppt", "image"}:
        # Show summary before table for PDF, PPT, images
        summary_text = response.content_summary or response.summary
        pre_table_content = f"""
        <div class='section'>
            <h3>📝 Summary</h3>
            <div class='content-box'>{html.escape(summary_text)}</div>
        </div>"""
    
    def display_visual_description(value: Optional[str]) -> str:
        cleaned = clean_text(value or "").strip(" .")
        normalized = " ".join(cleaned.lower().split())
        weak_values = {
            "ad",
            "advertisement",
            "diet",
            "food",
            "graphic",
            "health",
            "image",
            "n/a",
            "na",
            "no",
            "none",
            "no idea",
            "nothing",
            "photo",
            "picture",
            "poster",
            "product",
            "text",
            "unclear",
            "unknown",
            "words",
        }
        if not normalized or normalized in weak_values:
            return "N/A"
        if "no idea" in normalized or "do not know" in normalized or "don't know" in normalized:
            return "N/A"
        if len(normalized.split()) <= 2 and not any(
            term in normalized
            for term in ("person", "people", "package", "packaging", "bottle", "tablet", "food", "plate", "fork", "product")
        ):
            return "N/A"
        return cleaned

    rows = ""
    for match in response.key_message_matches:
        rows += f"""
        <tr>
          <td>{html.escape(match.key_message_id)}</td>
          <td>{html.escape(match.brand_product)}</td>
          <td>{html.escape(match.key_message)}</td>
          <td>{html.escape(match.key_words)}</td>
          <td>{html.escape(str(match.confidence_score))}</td>
          <td>{html.escape(match.source_location)}</td>
          <td>{html.escape(display_visual_description(match.description_of_image_video))}</td>
        </tr>
        """
    if not rows:
        rows = "<tr><td colspan='7'>No key-message matches found.</td></tr>"

    content = f"""
    <h1>✅ Intelligence Result</h1>
    <p><a href="/">← Analyze another file</a></p>
    <div class="notice">
      <b>Database Result ID:</b> {html.escape(str(response.metadata.get("analysis_result_id", "N/A")))}<br>
      <b>File:</b> {html.escape(response.file_name)}<br>
      <b>Type:</b> {html.escape(response.file_type)}<br>
      <b>Detected Language:</b> {html.escape(response.detected_language)}<br>
      <b>Target Language:</b> {html.escape(response.target_language)}
    </div>
    {warning_block}
    {pre_table_content}
    <h2>📊 Key Message Mapping</h2>
    <table>
      <thead>
        <tr>
          <th>Key Message ID</th>
          <th>Brand/Product</th>
          <th>Key Message</th>
          <th>Key Words</th>
          <th>Confidence</th>
          <th>Source</th>
          <th>Description of Image/Video</th>
        </tr>
      </thead>
      <tbody>{rows}</tbody>
    </table>
    """
    return HTMLResponse(_layout(content))


async def _run_analysis(
    file: UploadFile,
    key_message_file: Optional[UploadFile],
    target_language: str,
    mapping_enabled: bool,
    page_slide_number: Optional[int],
    top_k: Optional[int],
    ground_truth: Optional[UploadFile] = None,
    reference_transcript: Optional[UploadFile] = None,
) -> BaseIntelligenceResponse:
    uploaded_path = await save_upload_file(file, UPLOAD_DIR)
    upload_id = uploaded_path.stem
    original_filename = file.filename or uploaded_path.name
    await _save_optional_reference_file(upload_id, ground_truth, "ground_truth", {".txt"})
    await _save_optional_reference_file(upload_id, reference_transcript, "reference_transcript", {".txt", ".srt"})
    UPLOADED_MATERIALS[upload_id] = (uploaded_path, original_filename)
    UPLOADED_FILE_HASHES[upload_id] = _file_hash(uploaded_path)
    UPLOAD_IDS_BY_HASH.setdefault(UPLOADED_FILE_HASHES[upload_id], upload_id)
    return await _run_analysis_from_path(
        uploaded_path=uploaded_path,
        original_filename=original_filename,
        key_message_file=key_message_file,
        target_language=target_language,
        mapping_enabled=mapping_enabled,
        page_slide_number=page_slide_number,
        top_k=top_k,
    )


async def _run_analysis_from_path(
    uploaded_path: Path,
    original_filename: str,
    key_message_file: Optional[UploadFile],
    target_language: str,
    mapping_enabled: bool,
    page_slide_number: Optional[int],
    top_k: Optional[int],
) -> BaseIntelligenceResponse:

    if key_message_file and key_message_file.filename:
        key_message_path = await save_upload_file(key_message_file, UPLOAD_DIR)
    else:
        key_message_path = _latest_key_message_path()
        if not key_message_path:
            raise HTTPException(status_code=400, detail="Upload a key-message Excel/CSV file or place a key-message workbook in the project root.")

    analysis_path = uploaded_path
    file_hash = _upload_hash(uploaded_path.stem, uploaded_path)
    key_hash = _file_hash(key_message_path) if key_message_path and key_message_path.exists() else ""
    reference_dir = UPLOAD_DIR / uploaded_path.stem
    reference_hash = ""
    for reference_name in ("ground_truth.txt", "reference_transcript.txt", "reference_transcript.srt"):
        reference_path = reference_dir / reference_name
        if reference_path.exists():
            reference_hash += f"{reference_name}:{_file_hash(reference_path)};"
    analysis_cache_key = (
        "ocr_easyocr_path_v2",
        file_hash,
        key_hash,
        reference_hash,
        target_language or DEFAULT_TARGET_LANGUAGE,
        bool(mapping_enabled),
        page_slide_number,
        DEFAULT_TOP_K if top_k is None else top_k,
    )
    cached_payload = ANALYSIS_CACHE.get(analysis_cache_key)
    if cached_payload is not None:
        payload = _cache_copy(cached_payload)
        payload.setdefault("metadata", {})["cached"] = True
        response = BaseIntelligenceResponse(**payload)
        accuracy_report_path = save_accuracy_report(
            response,
            analysis_result_id=str(response.metadata.get("analysis_result_id") or ""),
        )
        response.metadata["accuracy_report_path"] = str(accuracy_report_path)
        return response

    response = analyze_file(
        file_path=analysis_path,
        original_filename=original_filename,
        key_message_path=key_message_path,
        target_language=target_language or DEFAULT_TARGET_LANGUAGE,
        mapping_enabled=mapping_enabled,
        page_slide_number=page_slide_number,
        top_k=DEFAULT_TOP_K if top_k is None else top_k,
    )
    response.metadata["upload_id"] = uploaded_path.stem
    response.metadata["original_filename"] = original_filename
    analysis_result_id = save_analysis_response(response)
    response.metadata["analysis_result_id"] = analysis_result_id
    try:
        response.metadata["mlr_review"] = run_mlr_review_for_analysis(
            uploaded_path.stem,
            original_filename,
            _response_payload(response),
        )
    except Exception as exc:
        response.metadata["mlr_review_error"] = str(exc)
    response.metadata["cached"] = False
    accuracy_report_path = save_accuracy_report(response, analysis_result_id=analysis_result_id)
    response.metadata["accuracy_report_path"] = str(accuracy_report_path)
    ANALYSIS_CACHE[analysis_cache_key] = _response_payload(response)
    return response


@app.post("/api/upload-material")
async def upload_material_api(
    file: UploadFile = File(...),
    ground_truth: Optional[UploadFile] = File(None),
    reference_transcript: Optional[UploadFile] = File(None),
):
    try:
        uploaded_path = await save_upload_file(file, UPLOAD_DIR)
        upload_id = uploaded_path.stem
        original_filename = _display_upload_name(uploaded_path)
        await _save_optional_reference_file(upload_id, ground_truth, "ground_truth", {".txt"})
        await _save_optional_reference_file(upload_id, reference_transcript, "reference_transcript", {".txt", ".srt"})
        file_hash = _file_hash(uploaded_path)
        UPLOADED_MATERIALS[upload_id] = (uploaded_path, original_filename)
        UPLOADED_FILE_HASHES[upload_id] = file_hash
        UPLOAD_IDS_BY_HASH.setdefault(file_hash, upload_id)
        preview = _upload_preview_metadata(upload_id, uploaded_path, original_filename)
        file_type = detect_file_type(original_filename)
        duplicate_upload_id = UPLOAD_IDS_BY_HASH.get(file_hash, upload_id)
        return {
            "upload_id": upload_id,
            "file_name": original_filename,
            "file_hash": file_hash,
            "file_type": file_type,
            "file_category": preview.get("category"),
            "preview": preview,
            "video_frames": [],
            "image_units": [],
            "cached_upload_id": duplicate_upload_id if duplicate_upload_id != upload_id else "",
            "processing_status": "Upload complete. Preview and analysis will run on demand.",
        }
    except ValueError as exc:
        return JSONResponse(status_code=400, content={"detail": str(exc)})
    except Exception as exc:
        traceback.print_exc()
        return JSONResponse(status_code=500, content={"detail": str(exc)})


@app.get("/api/file-types")
def file_types_api():
    return {
        "extensions": sorted(ALLOWED_EXTENSIONS),
        "types": {
            extension: {"extension": extension, **config}
            for extension, config in sorted(SUPPORTED_FILE_TYPES.items())
        },
    }


def _display_upload_name(path: Path) -> str:
    parts = path.name.split("_", 1)
    return parts[1] if len(parts) == 2 and len(parts[0]) >= 16 else path.name


async def _save_optional_reference_file(upload_id: str, upload_file: Optional[UploadFile], target_name: str, allowed_extensions: set[str]) -> None:
    if not upload_file or not upload_file.filename:
        return
    extension = Path(upload_file.filename).suffix.lower()
    if extension not in allowed_extensions:
        allowed = ", ".join(sorted(allowed_extensions))
        raise HTTPException(status_code=400, detail=f"{target_name} must be one of: {allowed}")
    target_dir = UPLOAD_DIR / upload_id
    target_dir.mkdir(parents=True, exist_ok=True)
    destination = target_dir / f"{target_name}{extension}"
    with destination.open("wb") as output:
        while True:
            chunk = await upload_file.read(1024 * 1024)
            if not chunk:
                break
            output.write(chunk)


@app.get("/api/uploaded-materials")
async def uploaded_materials_api():
    try:
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        items = []
        upload_paths: list[tuple[Path, float]] = []
        for path in UPLOAD_DIR.iterdir():
            try:
                stat = path.stat()
            except OSError:
                continue
            if not path.is_file():
                continue
            upload_paths.append((path, stat.st_mtime))
        for path, modified_at in sorted(upload_paths, key=lambda item: item[1], reverse=True):
            try:
                file_type = detect_file_type(path.name)
            except Exception:
                continue
            upload_id = path.stem
            original_filename = _display_upload_name(path)
            UPLOADED_MATERIALS[upload_id] = (path, original_filename)
            file_hash = _upload_hash(upload_id, path)
            items.append(
                {
                    "id": upload_id,
                    "name": original_filename,
                    "file_hash": file_hash,
                    "savedAt": datetime.fromtimestamp(modified_at).strftime("%d/%m/%Y, %H:%M:%S"),
                    "kind": file_type,
                    "category": get_file_type_config(path.name).get("category"),
                }
            )
        return {"documents": items}
    except Exception as exc:
        return JSONResponse(status_code=500, content={"detail": str(exc)})


@app.get("/api/materials/{upload_id}/preview")
def material_preview_api(upload_id: str):
    uploaded_path, original_filename = _get_uploaded_material(upload_id)
    file_hash = _upload_hash(upload_id, uploaded_path)
    cache_key = (file_hash, "metadata")
    cached = PREVIEW_METADATA_CACHE.get(cache_key)
    if cached is not None:
        return _preview_metadata_for_upload(cached, upload_id)
    preview = _upload_preview_metadata(upload_id, uploaded_path, original_filename)
    PREVIEW_METADATA_CACHE[cache_key] = _cache_copy(preview)
    return preview


@app.get("/api/material-preview/{upload_id}")
def material_preview_units_api(upload_id: str, start: int = 1, count: int = 1, quality: str = "thumb"):
    uploaded_path, original_filename = _get_uploaded_material(upload_id)
    file_type = detect_file_type(original_filename)
    file_hash = _upload_hash(upload_id, uploaded_path)
    safe_start = max(1, int(start or 1))
    max_preview_units = 500 if file_type in {"text", "script"} else 20
    safe_count = max(1, min(int(count or 1), max_preview_units))
    safe_quality = "full" if str(quality).lower() == "full" else "thumb"
    cache_key = (file_hash, file_type, safe_start, safe_count, safe_quality)
    cached = PREVIEW_UNITS_CACHE.get(cache_key)
    if cached is not None:
        return {"file_type": file_type, "units": _cache_copy(cached), "cached": True}
    if file_type == "pdf":
        units = _pdf_preview_units(uploaded_path, start=safe_start, count=safe_count, quality=safe_quality)
    elif file_type == "ppt":
        units = _powerpoint_preview_units(upload_id, uploaded_path, start=safe_start, count=safe_count, quality=safe_quality)
        if not units:
            converted_path = _convert_office_to_pdf(upload_id, uploaded_path)
            units = _pdf_preview_units(converted_path, label_prefix="Slide", start=safe_start, count=safe_count, quality=safe_quality) if converted_path and converted_path.exists() else []
        if not units:
            _, units, _ = _extract_preview_text(uploaded_path, original_filename)
            units = units[safe_start - 1 : safe_start - 1 + safe_count]
    elif file_type == "video":
        units = _video_preview_units(uploaded_path)
        units = units[safe_start - 1 : safe_start - 1 + safe_count]
    elif file_type == "image":
        units = _image_preview_units(uploaded_path)
        units = units[safe_start - 1 : safe_start - 1 + safe_count]
    elif file_type == "document" and uploaded_path.suffix.lower() in WORD_PREVIEW_EXTENSIONS:
        converted_path = _convert_office_to_pdf(upload_id, uploaded_path)
        units = _pdf_preview_units(converted_path, start=safe_start, count=safe_count, quality=safe_quality) if converted_path and converted_path.exists() else []
        if not units:
            _, units, _ = _extract_preview_text(uploaded_path, original_filename)
            units = units[safe_start - 1 : safe_start - 1 + safe_count]
    elif file_type in {"document", "spreadsheet", "email", "text", "script", "archive"}:
        _, units, _ = _extract_preview_text(uploaded_path, original_filename)
        units = units[safe_start - 1 : safe_start - 1 + safe_count]
    else:
        units = []
    PREVIEW_UNITS_CACHE[cache_key] = _cache_copy(units)
    return {"file_type": file_type, "units": units, "cached": False}


@app.get("/api/materials/{upload_id}/content")
def material_content_api(upload_id: str):
    uploaded_path, original_filename = _get_uploaded_material(upload_id)
    media_type = mimetypes.guess_type(original_filename)[0] or "application/octet-stream"
    return FileResponse(uploaded_path, media_type=media_type, filename=original_filename, content_disposition_type="inline")


@app.get("/api/materials/{upload_id}/preview-content")
def material_preview_content_api(upload_id: str):
    uploaded_path, original_filename = _get_uploaded_material(upload_id)
    converted_path = _converted_html_path(upload_id, uploaded_path)
    if not converted_path.exists():
        converted_path, _, _, _ = _convert_material_to_html_preview(upload_id, uploaded_path, original_filename)
    if not converted_path or not converted_path.exists():
        raise HTTPException(status_code=404, detail="Preview is not available for this file.")
    return FileResponse(converted_path, media_type="text/html; charset=utf-8", filename=converted_path.name, content_disposition_type="inline")


@app.get("/api/materials/{upload_id}/download")
def material_download_api(upload_id: str):
    uploaded_path, original_filename = _get_uploaded_material(upload_id)
    media_type = mimetypes.guess_type(original_filename)[0] or "application/octet-stream"
    return FileResponse(uploaded_path, media_type=media_type, filename=original_filename, content_disposition_type="attachment")


@app.get("/api/review-comments")
def list_review_comments(asset_id: str, content_version_id: Optional[str] = None, file_hash: Optional[str] = None):
    query: dict[str, Any] = {"asset_id": asset_id}
    file_hash = str(file_hash or "").strip()
    if file_hash:
        query = {"$or": [{"asset_id": asset_id}, {"file_hash": file_hash}]}
    if content_version_id:
        query = {"$and": [query, {"$or": [{"content_version_id": content_version_id}, {"file_hash": file_hash}]}]} if file_hash else {**query, "content_version_id": content_version_id}
    comments = review_annotations_collection.find(query).sort("created_at", 1)
    return {"comments": [_json_annotation(comment) for comment in comments]}


@app.get("/api/review-comments-history")
def list_review_comments_history(limit: int = 100):
    safe_limit = max(1, min(int(limit or 100), 500))
    comments = review_annotations_collection.find({}).sort("created_at", -1).limit(safe_limit)
    return {"comments": [_json_annotation(comment) for comment in comments]}


@app.post("/api/review-comments")
def create_review_comment(request: Request, payload: dict[str, Any] = Body(...)):
    current_user = _current_user_context(request)
    if not can_add_comment(current_user["role"]):
        raise HTTPException(status_code=403, detail="Your role cannot add review comments.")
    comment_text = str(payload.get("comment_text") or "").strip()
    asset_id = str(payload.get("asset_id") or "").strip()
    content_version_id = str(payload.get("content_version_id") or "").strip()
    file_hash = str(payload.get("file_hash") or "").strip()
    if not comment_text:
        raise HTTPException(status_code=400, detail="Comment text is required.")
    if not asset_id:
        raise HTTPException(status_code=400, detail="Asset id is required.")
    if not content_version_id:
        raise HTTPException(status_code=400, detail="Content version id is required.")

    annotation_type = str(payload.get("annotation_type") or "general").lower()
    if annotation_type not in {"general", "pin", "box", "timestamp", "copy_note"}:
        annotation_type = "general"
    status = str(payload.get("status") or "Open")
    if status not in {"Open", "Reopened", "Resolved", "Dismissed"}:
        status = "Open"
    department = current_user["department"]
    if is_admin(current_user["role"]):
        requested_department = str(payload.get("department") or "").strip()
        if requested_department in {"Admin", "Medical", "Legal", "Regulatory"}:
            department = requested_department
    now = datetime.utcnow()
    document = {
        "comment_text": comment_text,
        "file_name": str(payload.get("file_name") or "").strip(),
        "content_version_id": content_version_id,
        "asset_id": asset_id,
        "file_hash": file_hash,
        "annotation_type": annotation_type,
        "annotation_category": str(payload.get("annotation_category") or "general").strip() or "general",
        "copied_text": str(payload.get("copied_text") or payload.get("selected_text") or "").strip(),
        "page_number": payload.get("page_number"),
        "x": payload.get("x"),
        "y": payload.get("y"),
        "width": payload.get("width"),
        "height": payload.get("height"),
        "selection_rects": payload.get("selection_rects") if isinstance(payload.get("selection_rects"), list) else [],
        "timestamp": payload.get("timestamp"),
        "timestamp_end": payload.get("timestamp_end"),
        "mandatory": bool(payload.get("mandatory")),
        "status": status,
        "created_by_user": current_user["user_name"],
        "created_by_role": current_user["role"],
        "created_by_role_label": current_user["role_label"],
        "department": department,
        "created_at": now,
        "updated_at": now,
    }
    result = review_annotations_collection.insert_one(document)
    document["_id"] = result.inserted_id
    return _json_annotation(document)


def _response_payload(response: BaseIntelligenceResponse) -> dict[str, Any]:
    if hasattr(response, "model_dump"):
        return response.model_dump()
    return response.dict()


def _run_analysis_job(
    job_id: str,
    upload_id: str,
    target_language: str,
    mapping_enabled: bool,
    page_slide_number: Optional[int],
    top_k: Optional[int],
):
    ANALYSIS_JOBS[job_id] = {"status": "running", "created_at": datetime.utcnow().isoformat()}
    try:
        uploaded_path, original_filename = _get_uploaded_material(upload_id)
        response = asyncio.run(
            _run_analysis_from_path(
                uploaded_path,
                original_filename,
                None,
                target_language,
                mapping_enabled,
                page_slide_number,
                top_k,
            )
        )
        ANALYSIS_JOBS[job_id] = {
            "status": "completed",
            "created_at": ANALYSIS_JOBS[job_id].get("created_at"),
            "completed_at": datetime.utcnow().isoformat(),
            "result": _response_payload(response),
        }
    except Exception as exc:
        traceback.print_exc()
        ANALYSIS_JOBS[job_id] = {
            "status": "failed",
            "created_at": ANALYSIS_JOBS.get(job_id, {}).get("created_at", datetime.utcnow().isoformat()),
            "completed_at": datetime.utcnow().isoformat(),
            "detail": str(exc),
        }


@app.post("/api/analyze-jobs")
async def create_analysis_job(
    background_tasks: BackgroundTasks,
    upload_id: str = Form(...),
    target_language: str = Form(DEFAULT_TARGET_LANGUAGE),
    mapping_enabled: bool = Form(True),
    page_slide_number: Optional[int] = Form(None),
    top_k: Optional[int] = Form(DEFAULT_TOP_K),
):
    _get_uploaded_material(upload_id)
    job_id = uuid4().hex
    ANALYSIS_JOBS[job_id] = {"status": "queued", "created_at": datetime.utcnow().isoformat()}
    background_tasks.add_task(_run_analysis_job, job_id, upload_id, target_language, mapping_enabled, page_slide_number, top_k)
    return {"job_id": job_id, "status": "queued"}


@app.post("/api/analyze-jobs-json")
async def create_analysis_job_json(background_tasks: BackgroundTasks, payload: dict[str, Any] = Body(...)):
    upload_id = str(payload.get("upload_id") or "").strip()
    if not upload_id:
        raise HTTPException(status_code=400, detail="Uploaded file was not found. Please choose the file again.")
    target_language = str(payload.get("target_language") or DEFAULT_TARGET_LANGUAGE)
    mapping_enabled = bool(payload.get("mapping_enabled", True))
    page_slide_number = payload.get("page_slide_number")
    top_k = payload.get("top_k", DEFAULT_TOP_K)
    _get_uploaded_material(upload_id)
    job_id = uuid4().hex
    ANALYSIS_JOBS[job_id] = {"status": "queued", "upload_id": upload_id, "created_at": datetime.utcnow().isoformat()}
    background_tasks.add_task(_run_analysis_job, job_id, upload_id, target_language, mapping_enabled, page_slide_number, top_k)
    return {"job_id": job_id, "status": "queued"}


@app.get("/api/analyze-jobs/{job_id}")
def get_analysis_job(job_id: str):
    job = ANALYSIS_JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Analysis job was not found.")
    return job


@app.post("/api/mlr-review")
def mlr_review_api(payload: dict[str, Any] = Body(...)):
    try:
        result = payload.get("result") if isinstance(payload.get("result"), dict) else payload
        upload_id = str(payload.get("upload_id") or (result.get("metadata") or {}).get("upload_id") or "").strip()
        if upload_id:
            return run_mlr_review_for_analysis(upload_id, str(result.get("file_name") or upload_id), result)
        return run_mlr_review(payload)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"MLR review failed: {exc}") from exc


@app.patch("/api/review-comments/{comment_id}")
def update_review_comment(request: Request, comment_id: str, payload: dict[str, Any] = Body(...)):
    try:
        object_id = ObjectId(comment_id)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Invalid comment id.") from exc
    existing = review_annotations_collection.find_one({"_id": object_id})
    if not existing:
        raise HTTPException(status_code=404, detail="Comment was not found.")
    current_user = _current_user_context(request)
    allowed_fields = {"comment_text", "mandatory", "status"}
    update = {field: payload[field] for field in allowed_fields if field in payload}
    content_fields = {"comment_text", "mandatory"} & set(update)
    if content_fields and not can_edit_comment(
        current_user["role"],
        existing.get("created_by_role"),
        existing.get("created_by_user"),
        current_user["user_name"],
        _comment_department(existing),
    ):
        raise HTTPException(status_code=403, detail="You can edit only your own department comments.")
    if "status" in update and not can_resolve_comment(current_user["role"], _comment_department(existing)):
        raise HTTPException(status_code=403, detail="You can resolve or reopen only your own department comments.")
    if "status" in update and update["status"] not in {"Open", "Reopened", "Resolved", "Dismissed"}:
        raise HTTPException(status_code=400, detail="Invalid status.")
    if "comment_text" in update:
        update["comment_text"] = str(update["comment_text"]).strip()
    update["updated_at"] = datetime.utcnow()
    document = review_annotations_collection.find_one_and_update(
        {"_id": object_id},
        {"$set": update},
        return_document=ReturnDocument.AFTER,
    )
    if not document:
        raise HTTPException(status_code=404, detail="Comment was not found.")
    return _json_annotation(document)


@app.delete("/api/review-comments/{comment_id}")
def delete_review_comment(request: Request, comment_id: str):
    current_user = _current_user_context(request)
    if not can_delete_comment(current_user["role"]):
        raise HTTPException(status_code=403, detail="Only Admin can delete comments.")
    try:
        object_id = ObjectId(comment_id)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Invalid comment id.") from exc
    result = review_annotations_collection.delete_one({"_id": object_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Comment was not found.")
    return {"deleted": True, "id": comment_id}


@app.post("/analyze", response_class=HTMLResponse)
async def analyze_ui(
    file: UploadFile = File(...),
    key_message_file: Optional[UploadFile] = File(None),
    target_language: str = Form(DEFAULT_TARGET_LANGUAGE),
    mapping_enabled: bool = Form(False),
    page_slide_number: Optional[int] = Form(None),
    top_k: Optional[int] = Form(DEFAULT_TOP_K),
):
    try:
        response = await _run_analysis(file, key_message_file, target_language, mapping_enabled, page_slide_number, top_k)
        return _render_result(response)
    except Exception as exc:
        return HTMLResponse(_layout(f"<h1>Analysis Failed</h1><div class='error'>{html.escape(str(exc))}</div><p><a href='/'>Go back</a></p>"), status_code=500)


@app.post("/api/analyze", response_model=BaseIntelligenceResponse)
async def analyze_api(
    file: Optional[UploadFile] = File(None),
    upload_id: Optional[str] = Form(None),
    key_message_file: Optional[UploadFile] = File(None),
    ground_truth: Optional[UploadFile] = File(None),
    reference_transcript: Optional[UploadFile] = File(None),
    target_language: str = Form(DEFAULT_TARGET_LANGUAGE),
    mapping_enabled: bool = Form(True),
    page_slide_number: Optional[int] = Form(None),
    top_k: Optional[int] = Form(DEFAULT_TOP_K),
):
    try:
        if upload_id:
            uploaded = UPLOADED_MATERIALS.get(upload_id)
            if not uploaded:
                raise HTTPException(status_code=400, detail="Uploaded file was not found. Please choose the file again.")
            uploaded_path, original_filename = uploaded
            if not uploaded_path.exists():
                raise HTTPException(status_code=400, detail="Uploaded file is no longer available. Please choose the file again.")
            return await _run_analysis_from_path(
                uploaded_path,
                original_filename,
                key_message_file,
                target_language,
                mapping_enabled,
                page_slide_number,
                top_k,
            )
        if not file:
            raise HTTPException(status_code=400, detail="Choose a material file before analysis.")
        return await _run_analysis(file, key_message_file, target_language, mapping_enabled, page_slide_number, top_k, ground_truth, reference_transcript)
    except HTTPException:
        raise
    except Exception as exc:
        return JSONResponse(status_code=500, content={"detail": str(exc)})
